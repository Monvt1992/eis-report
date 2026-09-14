"""Assemble generated chart PNGs into a PPTX using the user's template."""

from __future__ import annotations

from pathlib import Path
from typing import List, Optional

from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def _hex_to_rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


class PptxReportBuilder:
    """Builds one slide per chart, reusing the provided template's layout
    and stretching each picture to fill the full content area of the slide.
    """

    def __init__(self, template_path: Optional[str], margin_in: float = 0.4):
        if template_path and Path(template_path).exists():
            self.prs = Presentation(template_path)
        else:
            self.prs = Presentation()  # blank default template
            self.prs.slide_width = Inches(13.333)
            self.prs.slide_height = Inches(7.5)

        self.margin = Inches(margin_in)
        # Prefer a mostly-blank layout (commonly index 6 "Blank", else the
        # last available layout) so charts aren't fighting template
        # placeholders for space.
        layouts = self.prs.slide_layouts
        self.chart_layout = layouts[6] if len(layouts) > 6 else layouts[-1]
        self.title_layout = layouts[0]

    def add_title_slide(self, title: str, subtitle: str = "") -> None:
        slide = self.prs.slides.add_slide(self.title_layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        if subtitle and len(slide.placeholders) > 1:
            try:
                slide.placeholders[1].text = subtitle
            except (KeyError, IndexError):
                pass

    def add_chart_slide(self, image_path: str | Path, heading: str,
                         milestone_title: Optional[str] = None,
                         milestone_color: str = "#C00000") -> None:
        slide = self.prs.slides.add_slide(self.chart_layout)

        sw, sh = self.prs.slide_width, self.prs.slide_height
        heading_h = Inches(0.6)

        # Heading text box.
        tb = slide.shapes.add_textbox(self.margin, Emu(0), sw - 2 * self.margin, heading_h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = heading
        run.font.size = Pt(22)
        run.font.bold = True

        if milestone_title:
            run2 = p.add_run()
            run2.text = f"   ({milestone_title})"
            run2.font.size = Pt(16)
            run2.font.bold = True
            run2.font.color.rgb = _hex_to_rgb(milestone_color)

        # Full-width/height content area below the heading.
        content_top = heading_h
        content_left = self.margin
        content_width = sw - 2 * self.margin
        content_height = sh - heading_h - self.margin

        slide.shapes.add_picture(
            str(image_path),
            left=content_left,
            top=content_top,
            width=content_width,
            height=content_height,
        )

    def save(self, out_path: str | Path) -> Path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(out_path)
        return out_path
